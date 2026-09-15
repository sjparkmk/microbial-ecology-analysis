import argparse, csv, hashlib
from pathlib import Path
from openpyxl import load_workbook

def sha(p: Path) -> str:
 h=hashlib.sha256();
 with p.open("rb") as f:
  for b in iter(lambda:f.read(1048576),b""): h.update(b)
 return h.hexdigest().upper()
def main():
 p=argparse.ArgumentParser();p.add_argument("--workbook",type=Path,default=Path("/input/youre/data/Supplementary_Tables_S1-S9_active_v3.xlsx"));p.add_argument("--output",type=Path,default=Path("/output/your/data/Table_S9_DNA_Library_Conc_v3.tsv"));p.add_argument("--log",type=Path,default=Path("/output/your/data/table_export_S9.log"));a=p.parse_args()
 if a.output.exists() or a.log.exists(): raise FileExistsError("Refusing to overwrite")
 wb=load_workbook(a.workbook,read_only=True,data_only=True);ws=wb["TableS9_DNA_Library_Conc"]
 lastrow=lastcol=0
 for ri,row in enumerate(ws.iter_rows(),1):
  for ci,c in enumerate(row,1):
   if c.value is not None:lastrow=max(lastrow,ri);lastcol=max(lastcol,ci)
 with a.output.open("w",encoding="utf-8",newline="") as h:
  w=csv.writer(h,delimiter="\t",lineterminator="\n")
  for row in ws.iter_rows(min_row=1,max_row=lastrow,max_col=lastcol,values_only=True):w.writerow(["" if x is None else x for x in row])
 text=a.output.read_text(encoding="utf-8")
 required=["normalized and pooled equimolarly before sequencing (author-confirmed)","normalization volumes were not separately itemized"]
 if not all(x in text for x in required):raise ValueError("Missing confirmed equimolar-pooling wording")
 a.log.write_text("PASS: Table S9 v3 TSV extraction\nworkbook="+str(a.workbook)+"\nworkbook_sha256="+sha(a.workbook)+"\noutput="+str(a.output)+"\noutput_sha256="+sha(a.output)+"\nconfirmed_pooling=normalized and pooled equimolarly before sequencing\n",encoding="utf-8")
 print("TABLE S9 V3: PASS")
if __name__=="__main__":main()
