from __future__ import annotations

import argparse
import csv
import hashlib
from pathlib import Path

from openpyxl import load_workbook


TABLES = {
    "TableS1_Alpha_Diversity": "Table_S1_Alpha_Diversity_v2.tsv",
    "TableS2_CPCoA_PERMDISP": "Table_S2_CPCoA_PERMDISP_v2.tsv",
    "TableS3_Domain_Composition": "Table_S3_Domain_Composition_v2.tsv",
    "TableS4_Phylum_Composition": "Table_S4_Phylum_Composition_v2.tsv",
    "TableS5_Niche_Breadth": "Table_S5_Niche_Breadth_v2.tsv",
    "TableS6_Read_Tracking": "Table_S6_Read_Tracking_v2.tsv",
    "TableS7_ASV_Bootstrap_Top30": "Table_S7_ASV_Bootstrap_Top30_v2.tsv",
    "TableS8_Environmental": "Table_S8_Environmental_v2.tsv",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def used_dimensions(ws) -> tuple[int, int]:
    last_row = 0
    last_column = 0
    for row_index, row in enumerate(ws.iter_rows(), start=1):
        for column_index, cell in enumerate(row, start=1):
            if cell.value is not None:
                last_row = row_index
                last_column = max(last_column, column_index)
    return last_row, last_column


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=Path("/input/youre/data/Supplementary_Tables_S1-S9_active_v3.xlsx"))
    parser.add_argument("--output-root", type=Path, default=Path("/output/your/data"))
    parser.add_argument("--log", type=Path, default=Path("/output/your/data/table_export_S1-S8.log"))
    args = parser.parse_args()
    if args.log.exists():
        raise FileExistsError(args.log)
    wb = load_workbook(args.workbook, read_only=True, data_only=True)
    rows_for_log = ["table_sheet\toutput\trows\tcolumns\tsha256"]
    for sheet, filename in TABLES.items():
        if sheet not in wb.sheetnames:
            raise KeyError(sheet)
        output = args.output_root / filename
        if output.exists():
            raise FileExistsError(output)
        ws = wb[sheet]
        rows, columns = used_dimensions(ws)
        with output.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
            for values in ws.iter_rows(min_row=1, max_row=rows, max_col=columns, values_only=True):
                writer.writerow(["" if value is None else value for value in values])
        rows_for_log.append(f"{sheet}\t{output.name}\t{rows}\t{columns}\t{sha256(output)}")
    s2 = (args.output_root / TABLES["TableS2_CPCoA_PERMDISP"]).read_text(encoding="utf-8")
    s7 = (args.output_root / TABLES["TableS7_ASV_Bootstrap_Top30"]).read_text(encoding="utf-8")
    if "5039" not in s2 or "999" not in s2:
        raise ValueError("Table S2 does not contain the confirmed CPCoA/PERMDISP permutation wording")
    if "Unclassified Gammaproteobacteria" not in s7 or "46" not in s7:
        raise ValueError("Table S7 lacks the required ASV_1 formal label or diagnostic bootstrap")
    args.log.parent.mkdir(parents=True, exist_ok=True)
    args.log.write_text(
        "Active supplementary table TSV extraction\n"
        f"workbook={args.workbook}\n"
        f"workbook_sha256={sha256(args.workbook)}\n"
        "validation=TableS2 contains 5,039 and 999; TableS7 contains Unclassified Gammaproteobacteria and 46\n"
        + "\n".join(rows_for_log)
        + "\n",
        encoding="utf-8",
    )
    print(f"EXPORTED {len(TABLES)} TSV tables")


if __name__ == "__main__":
    main()
