\
\
\
\
\


from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

try:
    from PIL import Image
except ImportError:
    Image = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None


SAMPLE_ORDER = [
    "1st_BFP",
    "1st_BFW1",
    "1st_BFW2",
    "2nd_BFP",
    "2nd_BFW",
    "2nd_wt",
    "3rd_wt",
]
TYPE_ORDER = ["P_biofilm", "W_biofilm", "water"]
TYPE_LABEL = {
    "P_biofilm": "AMD precipitate-associated biofilm",
    "W_biofilm": "wall-attached biofilm",
    "water": "water",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def serialize_tsv(rows: list[list[object]]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, delimiter="\t", lineterminator="\n")
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8")


def write_if_changed(path: Path, payload: bytes) -> str:
    if path.exists() and path.read_bytes() == payload:
        return "UNCHANGED"
    status = "UPDATED" if path.exists() else "CREATED"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)
    return status


def read_csv(path: Path, delimiter: str = ",") -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def valid(value: object) -> bool:
    return value is not None and str(value).strip() not in {"", "NA", "NaN", "nan", "None"}


def nearest_genus_label(tax: dict[str, str]) -> tuple[str, str]:
    if valid(tax.get("Genus")):
        return str(tax["Genus"]), "Assigned genus"
    for rank in ("Family", "Order", "Class", "Phylum", "Superphylum", "Kingdom"):
        if valid(tax.get(rank)):
            return f"Unclassified {tax[rank]}", f"Unclassified genus; displayed at {rank}"
    return "Unclassified", "Unclassified genus; no higher-rank label available"


def inventory_detail(path: Path) -> str:
    detail: dict[str, object] = {}
    suffix = path.suffix.lower()
    if Image is not None and suffix in {".png", ".tif", ".tiff", ".jpg", ".jpeg"}:
        with Image.open(path) as img:
            detail["pixels"] = [img.width, img.height]
            dpi = img.info.get("dpi")
            if dpi:
                detail["dpi"] = [round(float(dpi[0]), 3), round(float(dpi[1]), 3)]
    elif PdfReader is not None and suffix == ".pdf":
        reader = PdfReader(str(path))
        page = reader.pages[0]
        width_pt = float(page.mediabox.width)
        height_pt = float(page.mediabox.height)
        detail["page_mm"] = [round(width_pt * 25.4 / 72, 3), round(height_pt * 25.4 / 72, 3)]
        detail["pages"] = len(reader.pages)
    return json.dumps(detail, ensure_ascii=False, separators=(",", ":"))


def relative(path: Path, root: Path) -> str:
    try:
        return str(path.relative_to(root)).replace("\\", "/")
    except ValueError:
        return str(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, default=Path("/input/youre/data/"))
    parser.add_argument("--output-root", type=Path, default=Path("/output/your/data"))
    parser.add_argument("--workbook", type=Path, default=Path("/input/youre/data/Supplementary_Tables_S1-S9_active_v3.xlsx"))
    parser.add_argument("--reference-script", type=Path, default=Path("/input/youre/data/reference_tree_script.R"))
    parser.add_argument("--reference-pptx-dir", type=Path, default=Path("/input/youre/data/reference_layouts"))
    args = parser.parse_args()

    root = args.source_root.resolve()
    output_root = args.output_root.resolve()
    fig_dir = output_root / "figures"
    log_dir = output_root / "logs"
    source_dir = output_root / "source_data"
    package = root
    input_dir = package / "input_data"
    output_dir = package / "output_data"
    workbook_path = args.workbook.resolve()

    required = {
        "sample_sheet": input_dir / "sample_sheet.tsv",
        "cpcoa": input_dir / "primary_cpcoa_site_scores.tsv",
        "alpha": output_dir / "alpha_diversity_per_sample.csv",
        "seqtab": output_dir / "seqtab_final_long.csv",
        "taxa": output_dir / "taxa_silva144.csv",
        "ranked": output_dir / "asv_taxonomy_with_bootstrap_ranked.csv",
        "tree": output_dir / "asv_tree_all663.newick",
        "tree_tips": output_dir / "Fig4_circular_ASV_tree_tips.tsv",
        "workbook": workbook_path,
        "reference_script": args.reference_script.resolve(),
    }
    missing = [f"{name}: {path}" for name, path in required.items() if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing required inputs:\n" + "\n".join(missing))

    log_dir.mkdir(parents=True, exist_ok=True)
    source_dir.mkdir(parents=True, exist_ok=True)

    inventory_paths: list[tuple[str, Path]] = []
    for path in sorted(fig_dir.iterdir()):
        if path.is_file() and "_A4_" not in path.name:
            inventory_paths.append(("existing figure", path))
    for path in required.values():
        inventory_paths.append(("authoritative input", path))
    for name in ("fig2a.pptx", "fig2b.pptx", "fig3a.pptx", "fig3b.pptx", "fig4.pptx",
                 "sfigs2a_b.pptx", "sfigs3.pptx", "sfigs4a.pptx", "sfigs4b.pptx", "sfigs5_up.pptx"):
        path = args.reference_pptx_dir.resolve() / name
        if path.is_file():
            inventory_paths.append(("2024 layout reference", path))

    inventory_rows: list[list[object]] = [[
        "role", "path", "bytes", "sha256", "modified_utc", "format_detail"
    ]]
    for role, path in inventory_paths:
        stat = path.stat()
        inventory_rows.append([
            role,
            relative(path, root.parent.parent),
            stat.st_size,
            sha256(path),
            datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
            inventory_detail(path),
        ])
    inventory_status = write_if_changed(
        log_dir / "00_preflight_inventory_A4_v1.tsv", serialize_tsv(inventory_rows)
    )

    sample_sheet = read_csv(required["sample_sheet"], "\t")
    sample_meta = {row["analysis_name"]: row for row in sample_sheet}
    if list(sample_meta) != SAMPLE_ORDER:
        raise ValueError(f"Unexpected sample order: {list(sample_meta)}")
    for sample in SAMPLE_ORDER:
        if sample_meta[sample]["Type"] not in TYPE_ORDER:
            raise ValueError(f"Unexpected Type for {sample}: {sample_meta[sample]['Type']}")

    seq_rows = read_csv(required["seqtab"])
    tax_rows = read_csv(required["taxa"])
    ranked_rows = read_csv(required["ranked"])
    if len(seq_rows) != 663 or len(tax_rows) != 663 or len(ranked_rows) != 663:
        raise ValueError(
            f"Expected 663 ASVs; seqtab={len(seq_rows)}, taxa={len(tax_rows)}, ranked={len(ranked_rows)}"
        )
    if not all(sample in seq_rows[0] for sample in SAMPLE_ORDER):
        raise ValueError("seqtab_final_long.csv does not contain the expected seven sample columns")
    total_reads = sum(int(row[s]) for row in seq_rows for s in SAMPLE_ORDER)
    if total_reads != 716_915:
        raise ValueError(f"Unexpected final read total: {total_reads}")

    tax_by_sequence = {row["sequence"]: row for row in tax_rows}
    if len(tax_by_sequence) != 663:
        raise ValueError("Taxonomy sequence keys are not unique")
    for row in seq_rows:
        if row["sequence"] not in tax_by_sequence:
            raise ValueError(f"Missing taxonomy for {row['ASV_ID']}")

    generated: list[list[object]] = [["artifact", "status", "rows", "sha256"]]


    cpcoa = read_csv(required["cpcoa"], "\t")
    cpcoa_rows: list[list[object]] = [["Sample", "Type_code", "Type_label", "CAP1", "CAP2"]]
    for row in cpcoa:
        sample = row["Sample"]
        type_code = row["Type"]
        if sample_meta[sample]["Type"] != type_code:
            raise ValueError(f"CPCoA Type mismatch for {sample}")
        cpcoa_rows.append([sample, type_code, TYPE_LABEL[type_code], row["CAP1"], row["CAP2"]])
    path = source_dir / "Fig2_CPCoA_source_A4_v1.tsv"
    status = write_if_changed(path, serialize_tsv(cpcoa_rows))
    generated.append([relative(path, root), status, len(cpcoa_rows) - 1, sha256(path)])


    sample_totals = {sample: sum(int(row[sample]) for row in seq_rows) for sample in SAMPLE_ORDER}
    genus_counts: dict[str, Counter[str]] = defaultdict(Counter)
    genus_status: dict[str, str] = {}
    for row in seq_rows:
        label, status_label = nearest_genus_label(tax_by_sequence[row["sequence"]])
        genus_status[label] = status_label
        for sample in SAMPLE_ORDER:
            genus_counts[label][sample] += int(row[sample])
    genus_mean = {
        label: sum(genus_counts[label][sample] / sample_totals[sample] * 100 for sample in SAMPLE_ORDER)
        / len(SAMPLE_ORDER)
        for label in genus_counts
    }
    top_genus = sorted(genus_mean, key=lambda label: (-genus_mean[label], label))[:18]
    expected_top = [
        "Unclassified Gammaproteobacteria", "Metallibacterium", "Leptospirillum",
        "Unclassified Bacillati", "Unclassified Thermoplasmatales",
        "Unclassified Burkholderiales", "Acidibacillus--other", "Cuniculiplasma",
        "Unclassified Tepidisphaeraceae", "Acidibacter",
        "Unclassified Thermoplasmata--other", "Unclassified Group 1.1c",
        "Aciditerrimonas", "Unclassified Acidobacteriaceae", "Thiomonas",
        "Unclassified Bacteria", "Acidithiobacillus", "Acidiphilium",
    ]
    if top_genus != expected_top:
        raise ValueError(f"Top-18 display taxa changed: {top_genus}")
    genus_rows: list[list[object]] = [[
        "Taxon", "Taxon_status", "rank_order", "sample", "Type_code", "Type_label",
        "reads", "relative_abundance_pct"
    ]]
    for rank, label in enumerate(top_genus + ["Others"], start=1):
        for sample in SAMPLE_ORDER:
            if label == "Others":
                reads = sample_totals[sample] - sum(genus_counts[t][sample] for t in top_genus)
                status_label = "Collapsed remainder"
            else:
                reads = genus_counts[label][sample]
                status_label = genus_status[label]
            genus_rows.append([
                label, status_label, rank, sample, sample_meta[sample]["Type"],
                sample_meta[sample]["Type_label"], reads, f"{reads / sample_totals[sample] * 100:.10f}"
            ])
    path = source_dir / "Fig3_Genus_Bubble_source_A4_v1.tsv"
    status = write_if_changed(path, serialize_tsv(genus_rows))
    generated.append([relative(path, root), status, len(genus_rows) - 1, sha256(path)])


    alpha = read_csv(required["alpha"])
    alpha_header = list(alpha[0]) + ["Type_code"]
    alpha_rows: list[list[object]] = [alpha_header]
    for row in alpha:
        sample = row["sample"]
        alpha_rows.append([row[column] for column in alpha[0]] + [sample_meta[sample]["Type"]])
    path = source_dir / "SFig1_AlphaDiversity_source_A4_v1.tsv"
    status = write_if_changed(path, serialize_tsv(alpha_rows))
    generated.append([relative(path, root), status, len(alpha_rows) - 1, sha256(path)])


    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)

    phylum_sheet = workbook["TableS4_Phylum_Composition"]
    phylum_header = [phylum_sheet.cell(4, col).value for col in range(1, 10)]
    phylum_data: list[list[object]] = []
    row_index = 5
    while valid(phylum_sheet.cell(row_index, 1).value):
        phylum_data.append([phylum_sheet.cell(row_index, col).value for col in range(1, 10)])
        row_index += 1
    if len(phylum_data) != 39:
        raise ValueError(f"Expected 39 phyla in Table S4, found {len(phylum_data)}")
    phylum_rows = [phylum_header] + phylum_data
    path = source_dir / "TableS4_Phylum_Composition_A4_v1.tsv"
    status = write_if_changed(path, serialize_tsv(phylum_rows))
    generated.append([relative(path, root), status, len(phylum_data), sha256(path)])


    raw_phylum: dict[str, Counter[str]] = defaultdict(Counter)
    for row in seq_rows:
        tax = tax_by_sequence[row["sequence"]]
        phylum = tax["Phylum"] if valid(tax.get("Phylum")) else "Unclassified"
        for sample in SAMPLE_ORDER:
            raw_phylum[phylum][sample] += int(row[sample])
    workbook_phylum = {str(row[0]): row for row in phylum_data}
    if set(raw_phylum) != set(workbook_phylum):
        raise ValueError("Table S4 phylum labels do not match the raw-count aggregation")
    max_abs_phylum_diff = 0.0
    for phylum, counts in raw_phylum.items():
        wb_row = workbook_phylum[phylum]
        for sample_index, sample in enumerate(SAMPLE_ORDER, start=1):
            recomputed = counts[sample] / sample_totals[sample] * 100
            max_abs_phylum_diff = max(max_abs_phylum_diff, abs(recomputed - float(wb_row[sample_index])))
    if max_abs_phylum_diff > 1e-9:
        raise ValueError(f"Table S4 cross-check failed; max absolute difference={max_abs_phylum_diff}")

    phylum_mean = {str(row[0]): float(row[8]) for row in phylum_data}
    top_phyla = sorted(
        (p for p in phylum_mean if p != "Unclassified"),
        key=lambda p: (-phylum_mean[p], p),
    )[:8]
    expected_phyla = [
        "Pseudomonadota", "Thermoplasmatota", "Nitrospirota", "Bacillota--other",
        "Actinomycetota", "Planctomycetota", "Thermoproteota", "Acidobacteriota",
    ]
    if top_phyla != expected_phyla:
        raise ValueError(f"Top-eight assigned phyla changed: {top_phyla}")
    phylum_display_rows: list[list[object]] = [[
        "Phylum", "rank_order", "sample", "Type_code", "Type_label", "relative_abundance_pct"
    ]]
    display_phyla = top_phyla + ["Unclassified", "Others"]
    for rank, phylum in enumerate(display_phyla, start=1):
        for sample_index, sample in enumerate(SAMPLE_ORDER, start=1):
            if phylum == "Others":
                value = sum(
                    float(workbook_phylum[p][sample_index])
                    for p in workbook_phylum
                    if p not in top_phyla and p != "Unclassified"
                )
            else:
                value = float(workbook_phylum[phylum][sample_index])
            phylum_display_rows.append([
                phylum, rank, sample, sample_meta[sample]["Type"],
                sample_meta[sample]["Type_label"], f"{value:.10f}"
            ])
    for sample in SAMPLE_ORDER:
        total = sum(float(row[5]) for row in phylum_display_rows[1:] if row[2] == sample)
        if abs(total - 100.0) > 1e-7:
            raise ValueError(f"Displayed phylum stack does not sum to 100 for {sample}: {total}")
    path = source_dir / "SFig3_Phylum_Composition_display_A4_v1.tsv"
    status = write_if_changed(path, serialize_tsv(phylum_display_rows))
    generated.append([relative(path, root), status, len(phylum_display_rows) - 1, sha256(path)])

    niche_sheet = workbook["TableS5_Niche_Breadth"]

    niche_header = [niche_sheet.cell(8, col).value for col in range(1, 19)]
    niche_data = [[niche_sheet.cell(row, col).value for col in range(1, 19)] for row in range(9, 138)]
    if len(niche_data) != 129 or any(not valid(row[0]) for row in niche_data):
        raise ValueError("The primary Table S5 block is not exactly 129 populated genus rows")
    path = source_dir / "TableS5_Niche_Breadth_primary_A4_v1.tsv"
    status = write_if_changed(path, serialize_tsv([niche_header] + niche_data))
    generated.append([relative(path, root), status, len(niche_data), sha256(path)])
    if "classification" in niche_header or "eligible (baseline gate)" in niche_header:
        raise ValueError("Categorical niche fields are not permitted in the active v8 source table")


    top15_rows: list[list[object]] = [[
        "ASV_ID", "rank", "pct_of_total", "display_label", "label_status", "Genus_boot",
        "meets_minBoot80", "Kingdom", "Phylum"
    ]]
    asv_tax_by_id: dict[str, dict[str, str]] = {}
    for rank, row in enumerate(ranked_rows, start=1):
        asv_id = f"ASV_{rank}"
        tax = tax_by_sequence[row["ASV"]]
        asv_tax_by_id[asv_id] = tax
        if rank <= 15:
            label, label_status = nearest_genus_label(tax)
            boot = float(row["Genus_boot"])
            top15_rows.append([
                asv_id, rank, row["pct_of_total"], label, label_status, row["Genus_boot"],
                boot >= 80.0, tax["Kingdom"], tax["Phylum"]
            ])
    path = source_dir / "Fig4_Ferrovum_top15_source_A4_v1.tsv"
    status = write_if_changed(path, serialize_tsv(top15_rows))
    generated.append([relative(path, root), status, len(top15_rows) - 1, sha256(path)])


    tree_tip_names = set(re.findall(r"ASV_\d+", required["tree"].read_text(encoding="utf-8")))
    if len(tree_tip_names) != 663:
        raise ValueError(f"Expected 663 tip identifiers in the full tree, found {len(tree_tip_names)}")
    samples_by_type = {
        type_code: [sample for sample in SAMPLE_ORDER if sample_meta[sample]["Type"] == type_code]
        for type_code in TYPE_ORDER
    }
    reads_by_type = {
        type_code: sum(sample_totals[sample] for sample in samples)
        for type_code, samples in samples_by_type.items()
    }
    displayed_tree_rows = []
    for row in seq_rows:
        total_reads_asv = sum(int(row[sample]) for sample in SAMPLE_ORDER)
        pooled_pct = total_reads_asv / total_reads * 100.0
        if pooled_pct >= 0.2:
            displayed_tree_rows.append((row, total_reads_asv, pooled_pct))
    if len(displayed_tree_rows) != 56:
        raise ValueError(f"Expected 56 ASVs at pooled >=0.2%; found {len(displayed_tree_rows)}")
    tree_annotation: list[list[object]] = [[
        "OTUID", "Phylum_display", "P_biofilm", "W_biofilm", "water",
        "pct_pooled", "Genus_display", "Genus_boot"
    ]]
    for row, total_reads_asv, pooled_pct in sorted(displayed_tree_rows, key=lambda x: int(x[0]["ASV_ID"].split("_")[1])):
        asv_id = row["ASV_ID"]
        if asv_id not in tree_tip_names:
            raise ValueError(f"Displayed tip absent from tree: {asv_id}")
        tax = tax_by_sequence[row["sequence"]]
        phylum = tax["Phylum"] if valid(tax.get("Phylum")) else f"Unclassified {tax['Kingdom']}"
        genus_display, _ = nearest_genus_label(tax)
        tree_annotation.append([
            asv_id,
            phylum,
            f"{sum(int(row[sample]) for sample in samples_by_type['P_biofilm']) / reads_by_type['P_biofilm'] * 100:.10f}",
            f"{sum(int(row[sample]) for sample in samples_by_type['W_biofilm']) / reads_by_type['W_biofilm'] * 100:.10f}",
            f"{sum(int(row[sample]) for sample in samples_by_type['water']) / reads_by_type['water'] * 100:.10f}",
            f"{pooled_pct:.10f}",
            genus_display,
            ""
        ])
    path = source_dir / "Fig4_circular_ASV_tree_annotation_A4_v1.tsv"
    status = write_if_changed(path, serialize_tsv(tree_annotation))
    generated.append([relative(path, root), status, len(tree_annotation) - 1, sha256(path)])


    asv_type_sets: dict[str, set[str]] = {type_code: set() for type_code in TYPE_ORDER}
    samples_by_type = {
        type_code: [sample for sample in SAMPLE_ORDER if sample_meta[sample]["Type"] == type_code]
        for type_code in TYPE_ORDER
    }
    for row in seq_rows:
        for type_code, samples in samples_by_type.items():
            if sum(int(row[sample]) for sample in samples) > 0:
                asv_type_sets[type_code].add(row["ASV_ID"])
    p_set, w_set, water_set = (asv_type_sets[t] for t in TYPE_ORDER)
    venn = [
        ("P_biofilm_only", len(p_set - w_set - water_set)),
        ("W_biofilm_only", len(w_set - p_set - water_set)),
        ("water_only", len(water_set - p_set - w_set)),
        ("P_biofilm_and_W_biofilm_only", len((p_set & w_set) - water_set)),
        ("P_biofilm_and_water_only", len((p_set & water_set) - w_set)),
        ("W_biofilm_and_water_only", len((w_set & water_set) - p_set)),
        ("all_three_Types", len(p_set & w_set & water_set)),
    ]
    if [value for _, value in venn] != [102, 127, 133, 84, 24, 31, 162]:
        raise ValueError(f"Unexpected Venn region values: {venn}")
    if sum(value for _, value in venn) != 663:
        raise ValueError("Venn region union does not equal 663 ASVs")
    path = source_dir / "SFig10_Venn_regions_A4_v1.tsv"
    status = write_if_changed(path, serialize_tsv([["region", "ASV_count"]] + [list(x) for x in venn]))
    generated.append([relative(path, root), status, len(venn), sha256(path)])

    generated_path = log_dir / "02_data_preparation_artifacts_A4_v1.tsv"
    generated_status = write_if_changed(generated_path, serialize_tsv(generated))

    print(f"PRE-FLIGHT INVENTORY: {inventory_status}; {len(inventory_rows) - 1} files")
    print(f"RAW DATA CHECK: 663 ASVs; 7 samples; {total_reads:,} reads")
    print(f"PHYLUM CROSS-CHECK: 39 phyla; max absolute difference={max_abs_phylum_diff:.3e}")
    print(f"GENUS DISPLAY: {len(genus_counts)} labels; top 18 + Others fixed")
    print("NICHE PRIMARY: 129 named genera; continuous values only; categorical calls excluded")
    print(f"TREE DISPLAY: {len(displayed_tree_rows)} of 663 ASVs at pooled abundance >=0.2%")
    print(f"VENN REGIONS: {dict(venn)}; union={sum(value for _, value in venn)}")
    print(f"ARTIFACT REGISTER: {generated_status}; {generated_path}")
    print("DATA PREPARATION: PASS")


if __name__ == "__main__":
    main()
